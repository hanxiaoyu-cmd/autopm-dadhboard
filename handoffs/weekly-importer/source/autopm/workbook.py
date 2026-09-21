"""Read weekly reports without executing macros or modifying the source workbook.

The source coordinates, styles and project boundaries are retained as evidence.
Extraction is deliberately conservative: a template cell is never a project ID,
and a past date alone never means that a task has been completed.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from functools import lru_cache
import colorsys
import json
from pathlib import Path
import re
import warnings as pywarnings
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile

from .normalize import (
    clean_text,
    normalize_header,
    normalize_key,
    normalize_project_id,
    parse_date,
)

MAX_FILE_BYTES = 40 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_SHEET_CELLS = 1_500_000
MAX_SOURCE_CELLS = 80_000
MAX_TEXT_CHARS = 2_000_000
MAX_CELL_CHARS = 32_767

FIELD_ALIASES = {
    "project": "project_name",
    "projectname": "project_name",
    "sku": "sku",
    "skus": "sku",
    "region": "region",
    "category": "category",
    "subcategory": "sub_category",
    "status": "status",
    "pm": "pm",
    "npilead": "npi_lead",
    "xptlead": "npi_lead",
    "epm": "epm",
    "npdlead": "npd_lead",
    "tpm": "tpm",
    "pmo": "pmo",
    "type": "type",
    "capacity": "capacity",
    "brand": "brand",
    "quality": "quality",
    "qualitylead": "quality",
    "qe": "quality",
    "factory": "factory",
    "fty": "factory",
    "manufacturingfactory": "factory",
    "country": "country",
    "tooling": "tooling",
    "currentprogress": "current_progress",
    "engineeringremarks": "current_progress",
    "engupdatethisweek": "update_this_week",
    "updatethisweek": "update_this_week",
    "engineeringupdatethisweek": "update_this_week",
    "kickoffdate": "kick_off_date",
    "dqtpfinishdate": "dqtp_finish_date",
    "mpawdate": "mp_aw_date",
    "compliancecompletedate": "compliance_complete_date",
    "firstcrddate": "first_crd_date",
    "loa13wksplan": "loa_13weeks_plan",
}
JIRA_COUNT_FIELDS = {
    "total": "jira_total",
    "open": "jira_open",
    "verify": "jira_verify",
    "readytoclose": "jira_ready_to_close",
    "closed": "jira_closed",
}
SUPPLEMENTAL_FIELDS = frozenset(JIRA_COUNT_FIELDS.values()) | {"jira_link", "next_plm"}
CANONICAL_FIELDS = (
    frozenset(FIELD_ALIASES.values())
    | {"start_date", "mp_start_date"}
    | SUPPLEMENTAL_FIELDS
)
DATE_FIELDS = frozenset(k for k in CANONICAL_FIELDS if k.endswith("_date"))
PROJECT_LABELS = {"projectnumber", "projectno", "projectid", "projnumber"}
DATE_LABELS = {"update", "updateon", "updatedon", "reportdate", "weeklyreportdate"}
ISSUE_LABELS = {
    "keyissues",
    "keyissue",
    "issues",
    "keyissuestasks",
    "keyissuesandtasks",
}
EXCLUDED_LABELS = {
    "safety3rdpartytest",
    "safetyand3rdpartytest",
    "jir a summary".replace(" ", ""),
}
NULL_TEXT = {"", "na", "n/a", "none", "tbc", "tbd", "-", "--", "nil", "null"}
NS = {
    "s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


class WorkbookError(ValueError):
    """Source workbook cannot be interpreted safely."""


@lru_cache(maxsize=4096)
def field_key(text):
    """Allow one edit/transposition only if it identifies one canonical field.

    Exact known roles win first. A typo equally close to NPI and NPD is unknown.
    Short headings and identity/version labels are never fuzzy-matched.
    """
    key = normalize_header(text)
    if key in FIELD_ALIASES:
        return FIELD_ALIASES[key]
    if len(key) < 5 or key in PROJECT_LABELS | DATE_LABELS:
        return None

    def one_edit(other):
        if len(other) < 5 or abs(len(key) - len(other)) > 1:
            return False
        if len(key) == len(other):
            positions = [i for i in range(len(key)) if key[i] != other[i]]
            return len(positions) == 1 or (
                len(positions) == 2
                and positions[1] == positions[0] + 1
                and key[positions[0]] == other[positions[1]]
                and key[positions[1]] == other[positions[0]]
            )
        longer, shorter = (key, other) if len(key) > len(other) else (other, key)
        return any(longer[:i] + longer[i + 1 :] == shorter for i in range(len(longer)))

    matches = {target for alias, target in FIELD_ALIASES.items() if one_edit(alias)}
    return next(iter(matches)) if len(matches) == 1 else None


def _null(value):
    return clean_text(value).casefold() in NULL_TEXT


def _theme_colors(archive):
    try:
        root = ET.fromstring(archive.read("xl/theme/theme1.xml"))
        result = []
        for slot in root.find("a:themeElements/a:clrScheme", NS):
            color = next(iter(slot))
            result.append(color.get("lastClr") or color.get("val"))
        return result
    except (KeyError, TypeError, ET.ParseError):
        return []


def _color(color, theme):
    if color is None:
        return None
    kind = color.type
    value = getattr(color, kind, None)
    if not isinstance(value, (str, int, bool)):
        value = None
    result = {"type": kind, "value": value, "tint": float(color.tint or 0)}
    rgb = None
    if kind == "rgb" and isinstance(value, str):
        rgb = value[-6:]
    elif kind == "theme" and isinstance(value, int) and value < len(theme):
        rgb = theme[value]
    elif kind == "indexed" and isinstance(value, int):
        from openpyxl.styles.colors import COLOR_INDEXED

        if 0 <= value < len(COLOR_INDEXED):
            rgb = COLOR_INDEXED[value][-6:]
    if rgb and re.fullmatch(r"[0-9a-fA-F]{6}", rgb):
        # Excel tint is applied to HLS luminance, not independently to RGB.
        red, green, blue = (int(rgb[n : n + 2], 16) / 255 for n in (0, 2, 4))
        hue, light, saturation = colorsys.rgb_to_hls(red, green, blue)
        tint = result["tint"]
        light = light * (1 + tint) if tint < 0 else light * (1 - tint) + tint
        rgb = "".join(
            f"{round(v * 255):02X}" for v in colorsys.hls_to_rgb(hue, light, saturation)
        )
        result["rgb"] = rgb
    return result


def active_date_text(cell):
    """Discard wholly struck lines, never splice parts of a date together."""
    if cell.get("font", {}).get("strike"):
        return ""
    runs = cell.get("rich_text") or []
    if not cell.get("has_struck_text") and not any(r.get("strike") for r in runs):
        return cell["text"]
    if not runs or clean_text("".join(r["text"] for r in runs)) != clean_text(
        cell["text"]
    ):
        return ""
    lines = [[]]
    for run in runs:
        for char in run["text"].replace("\r\n", "\n").replace("\r", "\n"):
            if char == "\n":
                lines.append([])
            else:
                lines[-1].append((char, bool(run.get("strike"))))
    active = []
    for line in lines:
        styles = {strike for char, strike in line if not char.isspace()}
        if len(styles) > 1:
            return ""  # Partial strike cannot establish a replacement date.
        if styles == {False}:
            active.append("".join(char for char, _ in line))
    return "\n".join(active).strip()


def cell_source_date(cell, block):
    return source_date(active_date_text(cell), block)


def is_green_date(cell):
    """Only an explicit green date-cell fill establishes completion."""
    if not active_date_text(cell):
        return False
    fill = cell.get("fill", {})
    rgb = (fill.get("foreground") or {}).get("rgb")
    if fill.get("pattern") != "solid" or not rgb:
        return False
    red, green, blue = (int(rgb[n : n + 2], 16) / 255 for n in (0, 2, 4))
    hue, saturation, value = colorsys.rgb_to_hsv(red, green, blue)
    return 65 / 360 <= hue <= 165 / 360 and saturation >= 0.08 and value >= 0.25


def _rich_runs(node):
    result = []
    for run in node.findall("s:r", NS):
        strike = run.find("s:rPr/s:strike", NS)
        result.append(
            {
                "text": run.findtext("s:t", default="", namespaces=NS),
                "strike": strike is not None
                and strike.get("val", "1") not in {"0", "false"},
            }
        )
    return result if any(run["strike"] for run in result) else None


def _shared_rich_text(archive):
    try:
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return {}
    return {
        index: runs for index, item in enumerate(root) if (runs := _rich_runs(item))
    }


def _xml_metadata(archive, worksheet_path, shared_rich):
    root = ET.fromstring(archive.read(worksheet_path))
    merges = [n.attrib["ref"] for n in root.findall("s:mergeCells/s:mergeCell", NS)]
    formulas = {
        n.attrib["r"]: (n.findtext("s:f", namespaces=NS) or "[shared formula]")
        for n in root.findall("s:sheetData/s:row/s:c", NS)
        if n.find("s:f", NS) is not None
    }
    conditional = [
        n.get("sqref", "") for n in root.findall("s:conditionalFormatting", NS)
    ]
    hidden_rows = [
        int(n.get("r"))
        for n in root.findall("s:sheetData/s:row", NS)
        if n.get("hidden") == "1"
    ]
    rich = {}
    for cell in root.findall("s:sheetData/s:row/s:c", NS):
        runs = None
        if cell.get("t") == "s":
            value = cell.findtext("s:v", namespaces=NS)
            if value and value.isdigit():
                runs = shared_rich.get(int(value))
        elif cell.get("t") == "inlineStr":
            inline = cell.find("s:is", NS)
            if inline is not None:
                runs = _rich_runs(inline)
        if runs:
            rich[cell.get("r")] = runs
    return merges, formulas, conditional, hidden_rows, rich


def _right_value(cells, label, max_distance=8):
    if not is_eligible_cell(label):
        return None
    for cell in sorted(cells, key=lambda c: c["col"]):
        if (
            cell["row"] != label["row"]
            or not label["col"] < cell["col"] <= label["col"] + max_distance
        ):
            continue
        if cell.get("display_hidden"):
            continue
        if not is_eligible_cell(cell):
            return None
        key = normalize_header(cell["text"])
        if (
            field_key(cell["text"])
            or key in PROJECT_LABELS
            or key in DATE_LABELS
            or key in {"type", "capacity"}
        ):
            return None
        if not _null(cell["text"]) and cell.get("kind") != "error":
            return cell
        # Explicit N/A is a value placeholder, not permission to take the next label/value.
        return None
    return None


def _ref(sheet, coord):
    return f"'{sheet.replace(chr(39), chr(39) * 2)}'!{coord}"


def is_eligible_cell(cell):
    """Excel displays only the anchor value of a merged range.

    Some source XML retains older values in covered cells. Keep those cells in
    raw diagnostics, but never use them as identity, dates, text or model evidence.
    """
    address = cell.get("address") or cell.get("ref", "").rsplit("!", 1)[-1]
    return (
        not cell.get("display_hidden", False)
        and not cell.get("formula")
        and cell.get("kind") != "error"
        and (not cell.get("merged_anchor") or cell["merged_anchor"] == address)
    )


def source_date(value, block):
    return parse_date(
        value, block["report_date"], date_order=block.get("date_order", "AUTO")
    )


def read_workbook(path, report_date=None, *, date_order="AUTO"):
    """Return bounded cell evidence only from the unique Report worksheet.

    ``report_date`` fills genuinely missing block dates. Existing in-cell dates
    remain authoritative. The top-level date is the newest source date; callers
    MUST use each project's own date for freshness checks and version stamps.
    """
    path = Path(path)
    parse_date(
        None, date_order=date_order
    )  # Validate even when all cells are ISO dates.
    if not path.is_file():
        raise WorkbookError(f"找不到周报文件：{path.name}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise WorkbookError("请上传 .xlsx 或 .xlsm 周报；旧版 .xls 请先另存为 .xlsx。")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise WorkbookError("周报超过 40 MB 上限；请仅保留本次需要的周报工作表。")
    override = parse_date(report_date) if report_date else None
    if report_date and not override:
        raise WorkbookError("补充周报日期须为有效的 YYYY-MM-DD 日期。")
    warnings = []
    try:
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
    except ImportError as exc:
        raise WorkbookError("缺少 openpyxl，请运行安装脚本后重试。") from exc
    try:
        archive = ZipFile(path)
        if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
            archive.close()
            raise WorkbookError("Excel 解压后超过 200 MB 安全读取上限。")
        theme = _theme_colors(archive)
        shared_rich = _shared_rich_text(archive)
        with pywarnings.catch_warnings(record=True):
            pywarnings.simplefilter("always")
            workbook = openpyxl.load_workbook(
                path, read_only=True, data_only=True, keep_links=False
            )
    except (BadZipFile, OSError, ET.ParseError, KeyError, ValueError) as exc:
        if isinstance(exc, WorkbookError):
            raise
        raise WorkbookError(
            f"无法读取 Excel：{type(exc).__name__}。请在 Excel 中检查文件后另存为 .xlsx。"
        ) from exc
    projects, skipped, anchored_ids = [], [], []
    total_chars = total_cells = 0
    try:
        candidates = [
            sheet
            for sheet in workbook.worksheets
            if normalize_key(sheet.title) == "report"
        ]
        if not candidates:
            raise WorkbookError(
                "未找到 Report 工作表。仅支持解析 Report（忽略大小写和空格），请核对工作表名称。"
            )
        if len(candidates) != 1:
            raise WorkbookError(
                "多个工作表名称规范化后均为 Report，无法唯一确定周报主表；请重命名后再解析。"
            )
        if candidates[0].sheet_state != "visible":
            raise WorkbookError("Report 工作表已隐藏，请取消隐藏后再解析。")
        for sheet in workbook:
            if sheet not in candidates:
                skipped.append(
                    {"sheet": sheet.title, "reason": "非 Report 工作表，不参与解析"}
                )
        for sheet in candidates:
            if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_SHEET_CELLS:
                raise WorkbookError(
                    f"工作表 {sheet.title} 尺寸过大，未截断读取；请清理无用格式区域。"
                )
            merges, formulas, conditional, hidden_rows, rich = _xml_metadata(
                archive, sheet._worksheet_path, shared_rich
            )
            merge_rows = defaultdict(list)
            for merged in merges:
                c1, r1, c2, r2 = range_boundaries(merged)
                for row_number in range(r1, r2 + 1):
                    merge_rows[row_number].append(
                        (c1, c2, merged.split(":")[0], merged)
                    )
            cells = []
            with pywarnings.catch_warnings(record=True):
                pywarnings.simplefilter("always")
                for row in sheet.iter_rows():
                    for c in row:
                        if c.value is None:
                            continue
                        text = clean_text(c.value)
                        if not text:
                            continue
                        if len(text) > MAX_CELL_CHARS:
                            raise WorkbookError(
                                f"{sheet.title}!{c.coordinate} 文本超过上限，未截断。"
                            )
                        total_cells += 1
                        total_chars += len(text)
                        if (
                            total_cells > MAX_SOURCE_CELLS
                            or total_chars > MAX_TEXT_CHARS
                        ):
                            raise WorkbookError(
                                "周报内容超过读取上限，未截断；请分拆周报。"
                            )
                        cell = {
                            "ref": _ref(sheet.title, c.coordinate),
                            "address": c.coordinate,
                            "row": c.row,
                            "col": c.column,
                            "text": text,
                            "kind": "date"
                            if isinstance(c.value, (date, datetime))
                            else ("error" if c.data_type == "e" else "text"),
                            "number_format": c.number_format,
                            "hidden_row": c.row in hidden_rows,
                            "fill": {
                                "pattern": c.fill.patternType,
                                "foreground": _color(c.fill.fgColor, theme),
                            },
                            "font": {
                                "color": _color(c.font.color, theme),
                                "bold": bool(c.font.bold),
                                "strike": bool(c.font.strike),
                            },
                            "rich_text": rich.get(c.coordinate),
                            "has_struck_text": bool(c.font.strike)
                            or c.coordinate in rich,
                            "merged_anchor": c.coordinate,
                            "merged_range": None,
                        }
                        if c.coordinate in formulas:
                            cell["formula"] = formulas[c.coordinate]
                        for c1, c2, anchor, merged in merge_rows.get(c.row, []):
                            if c1 <= c.column <= c2:
                                cell["merged_anchor"], cell["merged_range"] = (
                                    anchor,
                                    merged,
                                )
                                break
                        if cell["merged_anchor"] != c.coordinate:
                            cell["display_hidden"] = True
                        cells.append(cell)
            anchors = [
                c
                for c in cells
                if is_eligible_cell(c) and normalize_header(c["text"]) in PROJECT_LABELS
            ]
            for c in cells:
                if c.get("formula") or c.get("kind") == "error":
                    warnings.append(
                        f"{c['ref']} 为公式缓存或 Excel 错误值，未用于自动提取；请核实后提供明确值。"
                    )
            if not anchors:
                skipped.append(
                    {"sheet": sheet.title, "reason": "未找到 Project Number 项目块"}
                )
                continue
            if conditional:
                warnings.append(
                    f"{sheet.title} 存在条件格式；完成状态只使用日期单元格的显式绿色填充。"
                )
            for i, anchor in enumerate(anchors):
                end = (
                    anchors[i + 1]["row"] - 1 if i + 1 < len(anchors) else sheet.max_row
                )
                block_cells = [c for c in cells if anchor["row"] <= c["row"] <= end]
                id_cell = _right_value(block_cells, anchor)
                if not id_cell:
                    value_labels = [
                        c
                        for c in block_cells
                        if normalize_header(c["text"]) in FIELD_ALIASES
                    ]
                    if any(_right_value(block_cells, c) for c in value_labels):
                        warnings.append(
                            f"{anchor['ref']} 项目块有内容但缺少 Project Number，已跳过。"
                        )
                    continue
                project_id = normalize_project_id(id_cell["text"])
                if not re.fullmatch(r"[A-Z0-9][A-Z0-9._/\-]{1,63}", project_id):
                    warnings.append(
                        f"{id_cell['ref']} 项目编号格式不明确，已跳过该项目块。"
                    )
                    continue
                anchored_ids.append(project_id)
                block_dates = []
                date_refs = []
                invalid_date = False
                for label in block_cells:
                    if not is_eligible_cell(label):
                        continue
                    if label["row"] > anchor["row"] + 3:
                        continue
                    if normalize_header(label["text"]) in DATE_LABELS:
                        value = _right_value(block_cells, label)
                        following = sorted(
                            (
                                c
                                for c in block_cells
                                if c["row"] == label["row"]
                                and label["col"] < c["col"] <= label["col"] + 8
                                and not c.get("display_hidden")
                            ),
                            key=lambda c: c["col"],
                        )
                        if following and (
                            following[0].get("formula")
                            or following[0].get("kind") == "error"
                            or following[0].get("has_struck_text")
                        ):
                            invalid_date = True
                        if value and not value.get("has_struck_text"):
                            parsed = parse_date(value["text"], date_order=date_order)
                            if parsed:
                                block_dates.append(parsed)
                                date_refs.append(value["ref"])
                            elif not _null(value["text"]):
                                invalid_date = True
                if invalid_date:
                    warnings.append(
                        f"{project_id} Update on 存在无法确定的日期，补充日期不能覆盖该冲突，已跳过。"
                    )
                    continue
                if len(set(block_dates)) > 1:
                    warnings.append(f"{project_id} 块内周报日期冲突，已跳过。")
                    continue
                block_date = block_dates[0] if block_dates else override
                if not block_date:
                    warnings.append(
                        f"{project_id} 缺少有效 Update on 日期，已跳过；可填写补充日期后重新解析。"
                    )
                    continue
                if not block_dates:
                    warnings.append(f"{project_id} 使用用户补充周报日期 {override}。")
                elif override and override != block_date:
                    warnings.append(
                        f"{project_id} 保留单元格日期 {block_date}；补充日期只用于缺失日期。"
                    )
                projects.append(
                    {
                        "project_id": project_id,
                        "project_id_source": id_cell["ref"],
                        "date_order": date_order,
                        "report_date": block_date,
                        "report_date_source": date_refs or ["user:report_date"],
                        "sheet": sheet.title,
                        "start_row": anchor["row"],
                        "end_row": end,
                        "source": _ref(
                            sheet.title,
                            f"{anchor['address']}:{openpyxl.utils.get_column_letter(sheet.max_column)}{end}",
                        ),
                        "cells": block_cells,
                    }
                )
    finally:
        workbook.close()
        archive.close()
    duplicate_ids = {key for key, count in Counter(anchored_ids).items() if count > 1}
    if duplicate_ids:
        warnings.append("重复项目编号保留各区块，后续按 SKU＋工厂核对：" + ", ".join(sorted(duplicate_ids)))
    if not projects:
        detail = "；".join(warnings[:8])
        raise WorkbookError(
            "没有可安全处理的项目。请检查 Project Number、Update on 和周报主表。"
            + detail
        )
    dates = sorted({p["report_date"] for p in projects})
    if len(dates) > 1:
        warnings.append(
            f"各项目更新日期不同（{dates[0]} 至 {dates[-1]}），同步时逐项目使用原始日期。"
        )
    file_dates = re.findall(r"(?<!\d)(20\d{6})(?!\d)", path.stem)
    if any(parse_date(d) and parse_date(d) not in dates for d in file_dates):
        warnings.append("文件名日期与单元格日期不同；采用各项目 Update on 日期。")
    return {
        "source": path.name,
        "report_date": dates[-1],
        "projects": projects,
        "warnings": warnings,
        "skipped_sheets": skipped,
        "schema_version": 1,
    }


def _put_field(project, key, value, refs, warnings):
    if value is None or _null(value):
        return
    if key in project.get("blocked_fields", []):
        return
    if key in project["fields"] and project["fields"][key] != value:
        warnings.append(
            f"{project['project_id']} {key} 存在多处冲突值，已排除该字段，等待核对。"
        )
        project.setdefault("blocked_fields", []).append(key)
        project["fields"].pop(key)
        project["evidence"]["fields"].pop(key, None)
        return
    project["fields"][key] = value
    project["evidence"]["fields"][key] = refs


def allowed_model_cells(block):
    """Expose only source-backed Jira fields; safety sections remain excluded."""
    cells = [c for c in block["cells"] if is_eligible_cell(c)]
    safety_start, sections = _supplemental_sections(block)
    allowed_refs = {
        entry[part]["ref"]
        for entry in supplemental_entries(block)
        for part in ("heading", "value", "context")
        if entry[part] is not None
    }
    result = []
    for cell in cells:
        if cell["row"] >= safety_start:
            continue
        if (
            any(
                section["row"] <= cell["row"] < end and cell["col"] >= section["col"]
                for section, end in sections
            )
            and cell["ref"] not in allowed_refs
        ):
            continue
        result.append(cell)
    return result


def _horizontal_span(cell):
    merged = cell.get("merged_range")
    if not merged:
        return cell["col"], cell["col"]

    def column(address):
        value = 0
        for character in re.sub(r"[^A-Za-z]", "", address).upper():
            value = value * 26 + ord(character) - ord("A") + 1
        return value

    endpoints = merged.split(":")
    return column(endpoints[0]), column(endpoints[-1])


def _vertical_end(cell):
    merged = cell.get("merged_range")
    return int(re.sub(r"[^0-9]", "", merged.split(":")[-1])) if merged else cell["row"]


def _supplemental_sections(block):
    cells = [c for c in block["cells"] if is_eligible_cell(c)]
    safety_rows = [
        c["row"]
        for c in cells
        if normalize_header(c["text"]) in EXCLUDED_LABELS - {"jirasummary"}
    ]
    safety_start = min(safety_rows or [block["end_row"] + 1])
    jira = [
        c
        for c in cells
        if normalize_header(c["text"]) == "jirasummary" and c["row"] < safety_start
    ]
    issue_rows = [
        c["row"] for c in cells if normalize_header(c["text"]) in ISSUE_LABELS
    ]
    sections = []
    for section in jira:
        boundaries = [
            r for r in issue_rows + [j["row"] for j in jira] if r > section["row"]
        ]
        end = min(boundaries + [safety_start, block["end_row"] + 1])
        sections.append((section, end))
    return safety_start, sections


def supplemental_field_value(entry):
    """Normalize a validated Jira/PLM entry without inventing or aggregating values."""
    return clean_text(entry["value"]["text"])


def _supplemental_entries(block, warnings):
    _, sections = _supplemental_sections(block)
    # Keep formula/error candidates for warnings, but never accept them as values.
    displayed = [
        c
        for c in block["cells"]
        if not c.get("display_hidden")
        and (
            not c.get("merged_anchor")
            or c["merged_anchor"]
            == (c.get("address") or c.get("ref", "").rsplit("!", 1)[-1])
        )
    ]
    entries = []
    for section, end in sections:
        left, right = _horizontal_span(section)
        # Older unmerged templates use the section label as the left boundary.
        if not section.get("merged_range"):
            right = max((c["col"] for c in displayed), default=left)
        cells = [
            c
            for c in displayed
            if _vertical_end(section) < c["row"] < end and left <= c["col"] <= right
        ]
        cell_at = {(c["row"], c["col"]): c for c in cells}
        for heading in cells:
            if not is_eligible_cell(heading):
                continue
            name = normalize_header(heading["text"])
            key = JIRA_COUNT_FIELDS.get(name) or {
                "jiralink": "jira_link",
                "nextplm": "next_plm",
            }.get(name)
            if not key:
                continue
            if key == "next_plm":
                heading_end = _horizontal_span(heading)[1]
                candidates = sorted(
                    (
                        c
                        for c in cells
                        if c["row"] == heading["row"]
                        and heading_end < c["col"] <= right
                    ),
                    key=lambda c: c["col"],
                )
                value = candidates[0] if candidates else None
                if value and (
                    normalize_header(value["text"]) in JIRA_COUNT_FIELDS
                    or normalize_header(value["text"]) in {"jiralink", "nextplm"}
                    or field_key(value["text"])
                ):
                    value = None
            else:
                value = cell_at.get((_vertical_end(heading) + 1, heading["col"]))
                if value and _horizontal_span(value) != _horizontal_span(heading):
                    warnings.append(
                        f"{value['ref']} {heading['text']} 标题与值的列范围不一致，未自动填入。"
                    )
                    value = None
            if value is None or _null(value["text"]):
                continue
            if _vertical_end(value) >= end or _horizontal_span(value)[1] > right:
                warnings.append(
                    f"{value['ref']} {heading['text']} 值超出 Jira 区域，未自动填入。"
                )
                continue
            if not is_eligible_cell(value):
                warnings.append(
                    f"{value['ref']} {heading['text']} 为公式缓存或 Excel 错误值，未自动填入。"
                )
                continue
            if value.get("has_struck_text") or value.get("font", {}).get("strike"):
                warnings.append(
                    f"{value['ref']} {heading['text']} 含删除线，未自动填入。"
                )
                continue
            text = clean_text(value["text"])
            if key in JIRA_COUNT_FIELDS.values() and not re.fullmatch(
                r"0|[1-9][0-9]*", text
            ):
                warnings.append(
                    f"{value['ref']} {heading['text']} 必须为明确的非负整数字面值，未自动填入。"
                )
                continue
            if key == "jira_link" and not re.fullmatch(
                r"https?://[^\s/]+(?:/[^\s]*)?", text, re.IGNORECASE
            ):
                warnings.append(
                    f"{value['ref']} Jira Link 缺少明确的 http(s) URL 文本，未自动填入。"
                )
                continue
            entries.append(
                {"key": key, "heading": heading, "value": value, "context": section}
            )
    return entries


def supplemental_entries(block):
    """Return valid Jira/PLM heading/value/context chains for rules and model checks.

    Generic Total/Open headings are interpreted only inside a Jira Summary
    section. Counts are literal nonnegative integers, including zero; Jira links
    require actual URL text. Neither formulas nor hyperlink display labels are
    read as values, and no external links are opened.
    """
    return _supplemental_entries(block, [])


def is_location_suffix(text):
    return bool(re.fullmatch(r"\(\s*[A-Za-z]{2,3}\s*\)", clean_text(text)))


def timeline_entries(block):
    """Pair timeline dates with their actual one/two-cell ordered heading chain."""
    cells = allowed_model_cells(block)
    bounds = [
        c["row"]
        for c in cells
        if normalize_header(c["text"]) in ISSUE_LABELS | {"currentprogress"}
    ]
    end = min(bounds or [block["start_row"] + 9])
    rows = defaultdict(list)
    cell_at = {(c["row"], c["col"]): c for c in cells}
    for cell in cells:
        if block["start_row"] + 3 < cell["row"] < end:
            rows[cell["row"]].append(cell)
    header_rows = [
        row
        for row, values in rows.items()
        if len(values) >= 3
        and sum(not parse_date(c["text"], block["report_date"]) for c in values) >= 3
        and not any(parse_date(c["text"], block["report_date"]) for c in values)
    ]
    entries, consumed = [], set()
    for row in sorted(header_rows):
        for label in rows[row]:
            if (
                label["ref"] in consumed
                or _null(label["text"])
                or is_location_suffix(label["text"])
            ):
                continue
            headings = [label]
            below = [
                cell_at[(r, label["col"])]
                for r in range(row + 1, min(row + 4, end))
                if (r, label["col"]) in cell_at
            ]
            value = below[0] if below else None
            if (
                value
                and value["row"] == row + 1
                and _horizontal_span(value) == _horizontal_span(label)
            ):
                continuation = is_location_suffix(value["text"]) or (
                    normalize_header(label["text"]) == "tra"
                    and clean_text(label["text"]).endswith("/")
                    and normalize_header(value["text"]) == "ecndd"
                )
                if continuation:
                    headings.append(value)
                    consumed.add(value["ref"])
                    value = below[1] if len(below) > 1 else None
                    if value and _horizontal_span(value) != _horizontal_span(label):
                        continue  # Different horizontal regions cannot form one heading/date chain.
            entries.append(
                {
                    "name": clean_text("\n".join(c["text"] for c in headings)),
                    "headings": headings,
                    "value": value,
                }
            )
    return entries


def parse_local(evidence):
    """Extract exact normalized headings, dated timeline cells, and issue rows."""
    warnings = list(evidence.get("warnings", []))
    projects = []
    for block in evidence["projects"]:
        cells = allowed_model_cells(block)
        cell_at = {(c["row"], c["col"]): c for c in cells}
        project = {
            "project_id": block["project_id"],
            "report_date": block["report_date"],
            "fields": {},
            "tasks": [],
            "issues": [],
            "source": block["source"],
            "evidence": {
                "project_id": [block["project_id_source"]],
                "report_date": block["report_date_source"],
                "fields": {},
            },
        }
        issue_headers = [
            c for c in cells if normalize_header(c["text"]) in ISSUE_LABELS
        ]
        current = [c for c in cells if normalize_header(c["text"]) == "currentprogress"]
        timeline_end = min(
            [c["row"] for c in current + issue_headers] or [block["start_row"] + 9]
        )
        brands = [
            c
            for c in cells
            if c["row"] == block["start_row"]
            and normalize_header(c["text"]) in {"shark", "ninja"}
        ]
        if len(brands) == 1:
            _put_field(
                project, "brand", brands[0]["text"], [brands[0]["ref"]], warnings
            )
        # Fields only appear before the timeline/current-progress sections.
        for label in cells:
            key = field_key(label["text"])
            if not key or (label["row"] >= timeline_end and key != "current_progress"):
                continue
            if key == "current_progress":
                value = cell_at.get((label["row"] + 1, label["col"]))
            else:
                value = _right_value(block["cells"], label)
            if value:
                content = (
                    cell_source_date(value, block)
                    if key in DATE_FIELDS
                    else clean_text(value["text"])
                )
                if key in DATE_FIELDS and content is None and not _null(value["text"]):
                    project.setdefault("blocked_fields", []).append(key)
                    project["fields"].pop(key, None)
                    warnings.append(
                        f"{value['ref']} {key} 日期有歧义或缺少明确年份，未自动填入。"
                    )
                _put_field(
                    project, key, content, [label["ref"], value["ref"]], warnings
                )
        for entry in _supplemental_entries(block, warnings):
            refs = [
                entry[part]["ref"]
                for part in ("context", "heading", "value")
                if entry[part] is not None
            ]
            _put_field(
                project, entry["key"], supplemental_field_value(entry), refs, warnings
            )
        # Reuse the same source heading chains in local and model validation.
        for entry in timeline_entries(block):
            name, value = entry["name"], entry["value"]
            if not value or _null(value["text"]):
                continue
            normalized = normalize_header(name)
            mp_source = normalized.startswith(("newmpstart", "mpstart"))
            if not active_date_text(value):
                if mp_source:
                    project.setdefault("blocked_fields", []).append("mp_start_date")
                    project["fields"].pop("mp_start_date", None)
                warnings.append(
                    f"{value['ref']} {name} 日期含删除线（含富文本），可能已作废，未自动写入。"
                )
                continue
            parsed = cell_source_date(value, block)
            if not parsed:
                if mp_source:
                    project.setdefault("blocked_fields", []).append("mp_start_date")
                    project["fields"].pop("mp_start_date", None)
                warnings.append(
                    f"{value['ref']} {name} 日期不明确：{value['text']!r}，未自动选择或猜测日期。"
                )
                continue
            if any(t["source"] == value["ref"] for t in project["tasks"]):
                continue
            refs = [c["ref"] for c in entry["headings"]] + [value["ref"]]
            project["tasks"].append(
                {
                    "name": name,
                    "date": parsed,
                    "completed": is_green_date(value),
                    "source": value["ref"],
                    "evidence": refs,
                }
            )
            normalized = normalize_header(name)
            key = {
                "award": "start_date",
                "kickoff": "kick_off_date",
                "dqtpfinish": "dqtp_finish_date",
                "mpaw": "mp_aw_date",
                "compliancecomplete": "compliance_complete_date",
            }.get(normalized)
            if normalized.startswith("newmpstart") or normalized.startswith("mpstart"):
                key = "mp_start_date"
            if key:
                _put_field(project, key, parsed, refs, warnings)
        for header in issue_headers:
            same_row = [c for c in cells if c["row"] == header["row"]]
            columns = {"text": header["col"]}
            for c in same_row:
                key = normalize_header(c["text"])
                field = {
                    "actions": "action",
                    "action": "action",
                    "owner": "owner",
                    "duedate": "due_date",
                    "riskhml": "risk",
                    "risk": "risk",
                }.get(key)
                if field:
                    columns[field] = c["col"]
            safety_rows = [
                c["row"]
                for c in block["cells"]
                if is_eligible_cell(c)
                and c["row"] > header["row"]
                and normalize_header(c["text"]) in EXCLUDED_LABELS
            ]
            next_headers = [c["row"] for c in issue_headers if c["row"] > header["row"]]
            end = min(safety_rows + next_headers or [block["end_row"] + 1])
            for row in sorted(
                {c["row"] for c in cells if header["row"] < c["row"] < end}
            ):
                text_cell = cell_at.get((row, columns["text"]))
                action_cell = cell_at.get((row, columns.get("action", -1)))
                if not text_cell or _null(text_cell["text"]):
                    if action_cell and not _null(action_cell["text"]):
                        warnings.append(
                            f"{action_cell['ref']} 只有 Action 没有 Issue，已保留在来源证据中等待核对。"
                        )
                    continue
                issue = {
                    "text": text_cell["text"],
                    "action": "",
                    "source": text_cell["ref"],
                    "evidence": [text_cell["ref"]],
                }
                for key, col in columns.items():
                    if key == "text":
                        continue
                    value = cell_at.get((row, col))
                    if value and not _null(value["text"]):
                        parsed = (
                            cell_source_date(value, block)
                            if key == "due_date"
                            else value["text"]
                        )
                        if parsed:
                            issue[key] = parsed
                            issue["evidence"].append(value["ref"])
                        else:
                            warnings.append(
                                f"{value['ref']} Issue 到期日不明确，未填入。"
                            )
                project["issues"].append(issue)
        projects.append(project)
    return {
        "source": evidence["source"],
        "report_date": evidence["report_date"],
        "projects": projects,
        "warnings": warnings,
        "parser": "local",
        "extraction_mode": "rules",
    }
