"""Synthetic tracker regression checks; no network or real business records."""
from datetime import date, datetime, time, timedelta
import hashlib
import json
from pathlib import Path
import re
import tempfile
import unittest
from zipfile import ZipFile

import openpyxl

from test_sync import fixture
from autopm.airtable_tracker import build_airtable_tracker_plan
from autopm.tracker import export_tracker, fingerprint, read_tracker, scalar
from autopm.tracker_roundtrip import writeback_plan
from autopm.workbook import WorkbookError


class TrackerTemporalValuesTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.source = self.folder / 'temporal-tracker.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'All Track'
        sheet.append(['Project Number', 'SKU', 'Manufacturing Factory',
                      'Project Name', 'Local clock', 'Elapsed time', 'Zero date',
                      'Report Date', 'FOT', 'Tooling', 'EB1'])
        sheet.append(['AB-123', 'SKU-1', 'F1', time(10, 5),
                      time(9, 30, 15, 123000), timedelta(days=2, hours=3, minutes=7),
                      0, '2026-09-01', None, time(8, 45, 1, 234000), 0])
        sheet['G2'].number_format = sheet['K2'].number_format = 'yyyy-mm-dd'
        # Orphan timeline cells must not prevent the valid project from updating.
        sheet['I4'] = time(0, 31, 0, 690000)
        sheet['I5'] = timedelta(days=2, hours=3)
        book.save(self.source)
        book.close()
        self.snapshot, _ = fixture()
        self.snapshot['captured_at'] = '2026-09-21T12:00:00+00:00'
        self.config = {'base_id': 'appTest', 'project_identity_mode': 'number_sku_factory'}
        project = self.snapshot['records']['tbl_projects'][0]
        project['fields'].update(fld_projects_sku='SKU-1',
                                 fld_projects_factory=['recFactory'],
                                 fld_projects_report_date='2026-09-20',
                                 fld_projects_project_name='Cloud demo name')

    def plan(self, path=None):
        return build_airtable_tracker_plan(path or self.source, self.snapshot, self.config)

    def cell_xml(self, path, coordinate):
        with ZipFile(path) as archive:
            xml = archive.read('xl/worksheets/sheet1.xml')
        pattern = rb'<(?:\w+:)?c\b[^>]*\br="' + coordinate.encode() + rb'"[^>]*>.*?</(?:\w+:)?c>'
        return re.search(pattern, xml, re.S)[0]

    def test_unmapped_time_zero_date_and_multiday_duration_are_json_safe(self):
        tracker = read_tracker(self.source)
        self.assertEqual(tracker['cells']['E2'], '09:30:15.123000')
        self.assertEqual(tracker['cells']['F2'], '2 days, 3:07:00')
        self.assertEqual(tracker['cells']['G2'], '00:00:00')
        json.dumps(tracker['cells'])
        self.assertEqual(tracker['content_sha256'], read_tracker(self.source)['content_sha256'])

    def test_scalar_preserves_subsecond_precision_without_inventing_dates(self):
        self.assertEqual(scalar(time(23, 59, 59, 123456)), '23:59:59.123456')
        self.assertEqual(scalar(timedelta(days=2, seconds=7, microseconds=123456)),
                         '2 days, 0:00:07.123456')
        self.assertEqual(scalar(timedelta(microseconds=-1)), '-1 day, 23:59:59.999999')

    def test_orphan_timeline_times_warn_without_blocking_the_valid_project(self):
        tracker = read_tracker(self.source)
        self.assertEqual(tracker['index'], {'AB-123': [2]})
        for coordinate, header in (('I4', 'FOT'), ('I5', 'FOT'), ('K2', 'EB1')):
            self.assertTrue(any('All Track' in warning and '!' + coordinate in warning
                                and header in warning for warning in tracker['warnings']))
        self.assertEqual(tracker['cells']['I4'], '00:31:00.690000')
        self.assertEqual(tracker['cells']['I5'], '2 days, 3:00:00')
        self.assertEqual(tracker['cells']['K2'], '00:00:00')
        plan = self.plan()
        self.assertEqual(plan['matched_projects'], 1)
        self.assertTrue(any(change['cell'] == 'D2' for change in plan['changes']))
        self.assertFalse(any(change['cell'] in {'I4', 'I5', 'K2'} for change in plan['changes']))
        for warning in tracker['warnings']:
            self.assertIn(warning, plan['warnings'])

    def test_roundtrip_preview_export_audit_and_unchanged_reread_preserve_temporal_cells(self):
        source_hash = fingerprint(self.source)
        plan = self.plan()
        change = next(item for item in plan['changes'] if item['semantic'] == 'project_name')
        self.assertEqual(change['before'], '10:05:00')
        self.assertEqual(change['after'], 'Cloud demo name')
        tooling = next(item for item in plan['roundtrip']['bindings'] if item['semantic'] == 'tooling')
        self.assertEqual(tooling['excel_before'], '08:45:01.234000')
        json.dumps(plan)
        output = self.folder / 'updated.xlsx'
        result = export_tracker(plan, output)
        audit = json.loads(Path(result['audit']).read_text(encoding='utf-8'))
        self.assertEqual(audit['changes'], plan['changes'])
        self.assertEqual(audit['roundtrip'], plan['roundtrip'])
        self.assertEqual(fingerprint(self.source), source_hash)
        for coordinate in ('E2', 'F2', 'G2', 'I4', 'I5', 'J2', 'K2'):
            self.assertEqual(self.cell_xml(self.source, coordinate), self.cell_xml(output, coordinate))
        book = openpyxl.load_workbook(output, read_only=True, data_only=False)
        try:
            self.assertEqual(book.active['D2'].value, 'Cloud demo name')
            for coordinate, expected in (
                    ('E2', time(9, 30, 15, 123000)),
                    ('F2', timedelta(days=2, hours=3, minutes=7)),
                    ('G2', time(0)), ('I4', time(0, 31, 0, 690000)),
                    ('I5', timedelta(days=2, hours=3)),
                    ('J2', time(8, 45, 1, 234000)), ('K2', time(0))):
                self.assertIsInstance(book.active[coordinate].value, type(expected))
                self.assertEqual(book.active[coordinate].value, expected)
        finally:
            book.close()
        repeated = self.plan(output)
        self.assertFalse(repeated['changes'])
        self.assertTrue(repeated['roundtrip']['bindings'])
        export_tracker(repeated, self.folder / 'unchanged.xlsx')
        writeback = writeback_plan(output, self.snapshot, self.config)
        self.assertFalse(writeback['blockers'])
        self.assertFalse(writeback['changes'])
        for warning in read_tracker(output)['warnings']:
            self.assertIn(warning, writeback['warnings'])
        json.dumps(writeback)

    def test_changed_time_or_duration_changes_content_hash_and_invalidates_legacy_audit(self):
        plan = self.plan()
        # Roundtrip baselines intentionally allow user edits; the legacy audit
        # must still reject a changed workbook instead of trusting stale hashes.
        plan.pop('roundtrip')
        for coordinate, replacement in (
                ('E2', time(9, 30, 15, 124000)),
                ('F2', timedelta(days=3, hours=3, minutes=7))):
            with self.subTest(coordinate=coordinate):
                output = self.folder / f'updated-{coordinate}.xlsx'
                export_tracker(plan, output)
                original_content_hash = read_tracker(output)['content_sha256']
                book = openpyxl.load_workbook(output)
                book.save(output)
                book.close()
                self.assertEqual(read_tracker(output)['content_sha256'], original_content_hash)
                self.assertFalse(self.plan(output)['changes'])
                book = openpyxl.load_workbook(output)
                book.active[coordinate] = replacement
                book.save(output)
                book.close()
                self.assertNotEqual(read_tracker(output)['content_sha256'], original_content_hash)
                with self.assertRaisesRegex(WorkbookError, '同步记录'):
                    self.plan(output)

    def test_cloud_calendar_date_replaces_time_only_format_on_target_cell(self):
        from openpyxl.styles.numbers import is_datetime
        self.snapshot['records']['tbl_tasks'] = [{'id': 'recFOT', 'fields': {
            'fld_tasks_name': 'FOT', 'fld_tasks_project': ['recProject'],
            'fld_tasks_milestone': 'FOT', 'fld_tasks_due_date': '2026-10-05'}}]
        for number_format, previous in (
                ('h:mm:ss', time(9, 30)),
                ('[h]:mm:ss', timedelta(days=2, hours=3))):
            with self.subTest(number_format=number_format):
                book = openpyxl.load_workbook(self.source)
                book.active['I2'] = previous
                book.active['I2'].number_format = number_format
                book.save(self.source)
                book.close()
                plan = self.plan()
                change = next(item for item in plan['changes'] if item['cell'] == 'I2')
                self.assertTrue(change['is_date'])
                self.assertEqual(change['after'], '2026-10-05')
                output = self.folder / ('calendar-from-' + type(previous).__name__ + '.xlsx')
                export_tracker(plan, output)
                for coordinate in ('I4', 'I5', 'E2', 'F2'):
                    self.assertEqual(self.cell_xml(self.source, coordinate), self.cell_xml(output, coordinate))
                book = openpyxl.load_workbook(output, read_only=True, data_only=False)
                try:
                    cell = book.active['I2']
                    self.assertIsInstance(cell.value, datetime)
                    self.assertEqual(cell.value.date(), date(2026, 10, 5))
                    self.assertIn(is_datetime(cell.number_format), {'date', 'datetime'})
                finally:
                    book.close()

    def test_existing_date_normalization_and_content_hash_remain_compatible(self):
        path = self.folder / 'date-only.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'All Track'
        sheet.append(['Project Number', 'Local date', 'Local datetime'])
        sheet.append(['AB-123', date(2026, 9, 19), datetime(2026, 9, 20, 13, 45)])
        book.save(path)
        book.close()
        tracker = read_tracker(path)
        expected = {
            'headers': {1: 'Project Number', 2: 'Local date', 3: 'Local datetime'},
            'index': {'AB-123': [2]},
            'values': {'A2': 'AB-123', 'B2': '2026-09-19', 'C2': '2026-09-20'},
        }
        legacy_hash = hashlib.sha256(json.dumps(expected, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
        self.assertEqual(tracker['cells'], expected['values'])
        self.assertEqual(tracker['content_sha256'], legacy_hash)


if __name__ == '__main__':
    unittest.main()
