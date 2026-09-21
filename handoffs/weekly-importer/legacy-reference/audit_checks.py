"""Read-only application audit. Synthetic workbooks and memory-only clients."""
import copy,json,sys,tempfile,socket
from pathlib import Path
from datetime import datetime
from unittest.mock import patch
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
from weekly_importer import WeeklyImporter,AllTrackerExporter,LedgerSync
from weekly_importer.remark_engine import read_weekly_remark,merge_weekly_remark
from weekly_importer.cleaner import clean_remark
import openpyxl
OUT=Path(__file__).resolve().parent
results=[]
class FakeClient:
    def __init__(self, rows):self.rows=copy.deepcopy(rows);self.writes=[]
    def get_records(self,table_id,fields=None,**kwargs):
        return [{'id':r['id'],'fields':{k:v for k,v in r['fields'].items() if fields is None or k in fields}} for r in self.rows]
    def patch_records(self,table_id,records):
        self.writes+=copy.deepcopy(records)
        for record in records:
            next(r for r in self.rows if r['id']==record['id'])['fields'].update(record['fields'])
        return {'records':copy.deepcopy(records)}
def row(rid,name='Example',remark='',**fields):
    return {'id':rid,'fields':{'Project name (Manual)':name,'Engineering remark(Manual)':remark,**fields}}
def workbook(path,rows,headers=('Date Added','Project Name , Description','Engineering Remarks')):
    w=openpyxl.Workbook();s=w.active;s.title='ALL PROJECTS'
    for i,h in enumerate(headers,1):s.cell(2,i,h)
    for r,data in enumerate(rows,3):
        for c,v in enumerate(data,1):s.cell(r,c,v)
    w.save(path);w.close()
def record(name,observed,detail):results.append({'check':name,'issue_reproduced':bool(observed),'detail':detail})
with patch.object(socket.socket,'connect',side_effect=AssertionError('Network disabled for offline audit')),tempfile.TemporaryDirectory() as temp:
    p=Path(temp)/'test.xlsx'
    workbook(p,[('2026-09-19','Example','New progress')])
    client=FakeClient([row('recFirst'),row('recLast')]);stats=WeeklyImporter(client,'Projects').import_from_excel(p)
    record('duplicate_project_names', [x['id'] for x in client.writes]==['recLast'], {'written_ids':[x['id'] for x in client.writes],'stats':stats})
    try:stats['status'];status_error='no error'
    except KeyError as exc:status_error=f'KeyError: {exc}'
    record('real_import_result_breaks_gui_cli_status_display',bool(client.writes) and 'status' not in stats,status_error)
    c=FakeClient([row('recOne')]);imp=WeeklyImporter(c,'Projects');imp.import_from_excel(p);c.writes=[];stats=imp.import_from_excel(p)
    results.append({'check':'simple_repeat_import','pass':not c.writes,'detail':stats})
    # Mixed order and conflicting same-name rows both derive from original remark.
    workbook(p,[('2026-09-18','Example','First progress'),('2026-09-19','Example','Second progress')])
    c=FakeClient([row('recOne')]);WeeklyImporter(c,'Projects').import_from_excel(p)
    record('duplicate_source_rows_same_target',len(c.writes)==2 and len({x['id'] for x in c.writes})==1, {'payload_count':len(c.writes),'first_update_lost_in_second_payload':'First progress' not in c.writes[-1]['fields']['Engineering remark(Manual)']})
    current='--- 2026-09-19 ---\nNewer progress'
    merged=merge_weekly_remark(current,'2026-09-01',{'current_progress':'Older progress'})
    record('stale_report_accepted',merged!=current,merged)
    same=merge_weekly_remark(current,'2026-09-20',{'current_progress':'Newer progress'})
    record('unchanged_content_loses_new_report_date',read_weekly_remark(same)['report_date']=='2026-09-19',read_weekly_remark(same))
    legacy='[[AUTOPM_WEEKLY_REPORT:v1]]\n周报日期: 2026-09-14\n工程进展:\nHistoric progress\n[[AUTOPM_WEEKLY_REPORT:END_PROGRESS]]\n[[/AUTOPM_WEEKLY_REPORT]]'
    record('actual_legacy_format_not_parsed',read_weekly_remark(legacy)['report_date'] is None,read_weekly_remark(legacy))
    cleaned=clean_remark('Manual note\n\n'+legacy)
    record('clean_removes_legacy_progress_content','Historic progress' not in cleaned,cleaned)
    workbook(p,[('Example','Progress')],headers=('Project Name , Description','Engineering Remarks'))
    try:WeeklyImporter(FakeClient([row('recOne')]),'Projects').import_from_excel(p,report_date='2026-09-19');err='no error'
    except Exception as exc:err=f'{type(exc).__name__}: {exc}'
    record('missing_date_header_crashes_even_with_override',err.startswith('ValueError'),err)
    client=FakeClient([row('recOne',remark=current,Status='On Track',Priority='High',Owner=['recPerson'],Notes='Manual note')])
    export=AllTrackerExporter(client,'Projects').export(p)
    syncstats=LedgerSync(client,'Projects').sync_from_excel(p)
    record('untouched_export_clears_cloud_fields',bool(client.writes) and client.writes[0]['fields'].get('Status','missing') is None,{'export':export,'sync':syncstats,'writes':client.writes})
    # A cloud edit after export cannot be distinguished from a local change.
    client=FakeClient([row('recOne',Status='New cloud value')]);w=openpyxl.load_workbook(p);w.active.cell(3,6,'Old local value');w.save(p);w.close()
    stats=LedgerSync(client,'Projects').sync_from_excel(p)
    record('cloud_conflict_not_detected',stats['conflicts']==0 and client.writes[0]['fields']['Status']=='Old local value',stats)
    client=FakeClient([row('recOne',remark=current)]);AllTrackerExporter(client,'Projects').export(p);WeeklyImporter(client,'Projects').import_from_excel(p)
    record('export_import_nests_full_history',bool(client.writes),client.writes)
    # No real credentials or cleaning requests: check CLI dispatch only.
    import weekly_importer.__main__ as cli
    fake_cleaner=type('C',(),{'clean_all':lambda self:{'total':0,'cleaned':0,'skipped':0,'errors':0}})()
    with patch('weekly_importer.load_config',return_value={'api_token':'synthetic','base_id':'appExample'}),patch('weekly_importer.AirtableClient'),patch('weekly_importer.RemarkCleaner',return_value=fake_cleaner) as clean,patch.object(sys,'argv',['audit','--clean','--dry-run']):
        cli.run_cli();record('clean_ignores_dry_run',clean.called,{'cleaner_dispatched':clean.called})
    # Real requests.Response truthiness reproduces failed retry selection, without HTTP.
    import requests
    import weekly_importer.api_client as api
    response=requests.Response();response.status_code=429
    with patch.object(api.requests,'request',return_value=response) as request,patch.object(api.time,'sleep'):
        try:api.AirtableClient('synthetic','appExample').get_records('Projects')
        except requests.exceptions.HTTPError:pass
        record('429_not_retried',request.call_count==1,{'calls':request.call_count,'configured_attempts':3})
(OUT/'离线复现结果.json').write_text(json.dumps(results,ensure_ascii=False,indent=2,default=str),encoding='utf8')
print(json.dumps({'checks':len(results),'issues_reproduced':sum(x.get('issue_reproduced',False) for x in results),'simple_repeat_pass':next(x['pass'] for x in results if x['check']=='simple_repeat_import')},ensure_ascii=False))
