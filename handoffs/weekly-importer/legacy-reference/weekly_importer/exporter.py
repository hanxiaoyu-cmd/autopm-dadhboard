"""Export Airtable Projects table to local Excel (All Tracker format)."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .api_client import AirtableClient


class AllTrackerExporter:
    """Export Airtable Projects data to standard All Tracker Excel format."""

    # Standard column headers matching the weekly report Excel format
    DEFAULT_HEADERS = [
        "Date Added",
        "Project Name , Description",
        "Engineering Remarks",
        "Project Number",
        "Project SKU",
        "Status",
        "Priority",
        "Owner",
        "Due Date",
        "Notes",
    ]

    # Hidden metadata column for round-trip sync tracking
    META_COL = 100

    def __init__(
        self,
        client: AirtableClient,
        table_id: str,
        project_name_field: str = "Project name (Manual)",
        remark_field: str = "Engineering remark(Manual)",
    ) -> None:
        self.client = client
        self.table_id = table_id
        self.project_name_field = project_name_field
        self.remark_field = remark_field

    def export(
        self,
        output_path: str | Path,
        *,
        include_fields: list[str] | None = None,
        max_records: int | None = None,
    ) -> dict[str, Any]:
        """Export Projects table to Excel.

        Args:
            output_path: Where to save the .xlsx file.
            include_fields: Which Airtable fields to fetch. Defaults to
                project name, remark, project number and SKU.
            max_records: Optional cap on number of records.

        Returns:
            Dict with exported count, path, and metadata.
        """
        fields = include_fields or [
            self.project_name_field,
            self.remark_field,
            "Project ID (Manual)",
            "Project SKU (Manual)",
        ]

        records = self.client.get_records(
            self.table_id,
            fields=fields,
            page_size=100,
        )

        if max_records:
            records = records[:max_records]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ALL PROJECTS"

        # Header row (row 2)
        for col_idx, header in enumerate(self.DEFAULT_HEADERS, start=1):
            cell = ws.cell(row=2, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="2563EB", end_color="2563EB", fill_type="solid"
            )
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Data rows (starting from row 3)
        for row_idx, record in enumerate(records, start=3):
            fields_data = record.get("fields", {})

            ws.cell(
                row=row_idx,
                column=1,
                value=datetime.now().strftime("%Y-%m-%d"),
            )
            ws.cell(
                row=row_idx,
                column=2,
                value=fields_data.get(self.project_name_field, ""),
            )
            ws.cell(
                row=row_idx,
                column=3,
                value=fields_data.get(self.remark_field, ""),
            )
            ws.cell(
                row=row_idx,
                column=4,
                value=fields_data.get("Project ID (Manual)", ""),
            )
            ws.cell(
                row=row_idx,
                column=5,
                value=fields_data.get("Project SKU (Manual)", ""),
            )

            # Hidden metadata for round-trip sync
            meta = {
                "record_id": record.get("id", ""),
                "exported_at": datetime.now().isoformat(),
            }
            ws.cell(row=row_idx, column=self.META_COL, value=json.dumps(meta))

        # Hide metadata column
        ws.column_dimensions[get_column_letter(self.META_COL)].hidden = True

        # Auto-adjust column widths
        for col_idx in range(1, len(self.DEFAULT_HEADERS) + 1):
            max_length = 0
            column = get_column_letter(col_idx)
            for cell in ws[column]:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 50)

        wb.save(output_path)

        return {
            "status": "exported",
            "count": len(records),
            "path": str(output_path),
            "fields": fields,
        }
